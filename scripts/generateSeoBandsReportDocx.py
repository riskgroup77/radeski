#!/usr/bin/env python3
"""Generate Word report: SEO bands 4-10 implementation analysis for Radeski.uz."""
from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "RADESKI_SEO_4_10_BANDLAR_TAHLILI.docx"

KOKAND_SLUGS = [
    "dermatolog", "trixolog", "podolog", "akne-davolash", "soch-tokilish",
    "tirnoq-zamburug", "osma-olib-tashlash", "psoriaz-davolash", "vitiligo-davolash",
    "rozasea-davolash", "pigmentatsiya", "lazer-epilyatsiya", "dermatoskopiya",
    "qon-tomir", "sogal-olib-tashlash",
]

FERGANA_SLUGS = [
    "dermatolog", "trixolog", "podolog", "akne-davolash", "psoriaz-davolash",
    "vitiligo-davolash", "rozasea-davolash", "soch-tokilish", "dermatoskopiya",
    "osma-olib-tashlash", "tirnoq-zamburug", "ipl-fargona", "pigmentatsiya",
    "lazer-tomir", "lazer-epilyatsiya",
]


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(item, style="List Bullet")
        p.paragraph_format.space_after = Pt(4)


def add_numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(item, style="List Number")
        p.paragraph_format.space_after = Pt(4)


def add_table_urls(doc: Document, city: str, slugs: list[str], queries_uz: list[str]) -> None:
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "№"
    hdr[1].text = "Qidiruv so'rovi (asosiy)"
    hdr[2].text = "URL (o'zbekcha)"
    for i, (slug, query) in enumerate(zip(slugs, queries_uz), start=1):
        row = table.add_row().cells
        row[0].text = str(i)
        row[1].text = query
        row[2].text = f"https://radeski.uz/uz/{city}/{slug}"


def main() -> None:
    doc = Document()
    title = doc.add_heading("RADESKI.UZ — SEO TAHLIL VA AMALGA OSHIRILGAN ISHLAR", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    sub = doc.add_paragraph(
        f"4–10-bandlar bo'yicha to'liq tahlil va implementatsiya hisoboti\n"
        f"Sanasi: {date.today().strftime('%d.%m.%Y')} | radeski.uz"
    )
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in sub.runs:
        run.font.size = Pt(11)
        run.font.italic = True

    doc.add_paragraph()
    doc.add_paragraph(
        "Ushbu hujjat 4dan10.docx hujjatidagi 4–10-bandlar (boblar) asosida "
        "radeski.uz saytida amalga oshirilgan SEO ishlarning batafsil tahlilini "
        "o'z ichiga oladi. Har bir bob alohida ko'rib chiqilgan."
    )

    # BOB 4
    add_heading(doc, "4-BOB. JORIY POZITSIONLASH (POSITIONING)", 1)
    doc.add_paragraph(
        "Manba tahlildagi tavsiya: Radeski «maxsus teri, soch va tirnoqlar klinikasi» "
        "sifatida pozitsiyalansin — oddiy kosmetologiya markazi emas."
    )
    add_heading(doc, "Amalga oshirilgan ishlar", 2)
    add_bullets(doc, [
        "POSITIONING_FORMULA — uch til (uz/ru/en) uchun yagona pozitsionlash formulasi katalogga qo'shildi.",
        "Qo'qon hub (/uz/qoqon) va Farg'ona hub (/uz/fargona) sahifalarida formula ko'rsatiladi.",
        "30 ta commercial landing sahifasida «Klinika haqida» blokida formula takrorlanadi.",
        "Home va About SEO meta allaqachon «teri, soch va tirnoqlar» ixtisoslashuvini aks ettiradi.",
        "Shahar nomlari endi alohida: dual «Farg'ona | Qo'qon» o'rniga har shahar uchun alohida hub va title.",
    ])
    add_heading(doc, "Formulasi (o'zbekcha)", 2)
    p = doc.add_paragraph()
    p.add_run(
        "«Teri, soch va tirnoqlar bo'yicha ixtisoslashgan klinika: dermatologiya, "
        "trixologiya, podologiya, diagnostika va zamonaviy apparat texnologiyalari.»"
    ).italic = True
    add_heading(doc, "Natija", 2)
    doc.add_paragraph(
        "Google va bemorlar uchun Radeski endi bir xil xabar bilan ko'rinadi: "
        "tibbiy ixtisoslashgan klinika. Bu 4-bandning asosiy strategik maqsadini qoplaydi."
    )

    # BOB 5
    add_heading(doc, "5-BOB. QO'QON RAQOBAT MUHITI", 1)
    doc.add_paragraph(
        "Tahlil: Qo'qonda Dermomeduz, Avitsenna, davlat KVD va alohida kosmetologlar "
        "raqobat qiladi. Radeski afzalligi — dermatologiya + trixologiya + podologiya + "
        "apparat + diagnostika bir brendda."
    )
    add_heading(doc, "Amalga oshirilgan ishlar", 2)
    add_bullets(doc, [
        "Qo'qon uchun alohida SEO hub: https://radeski.uz/uz/qoqon — filial manzili, telefon, xarita.",
        "15 ta maxsus commercial sahifa yaratildi — har biri bitta qidiruv niyatiga (intent) yo'naltirilgan.",
        "Raqobatchilardan farq qiluvchi 4 ta afzallik (6-band) Qo'qon hub va landinglarda ko'rsatiladi.",
        "Ichki linking: xizmatlar, maqolalar (psoriaz Daavlin Qo'qon, thulium lazer, onixokriptoz va b.) bog'landi.",
        "Footer va hub orqali Qo'qon sahifalariga doimiy ichki linklar.",
    ])
    add_heading(doc, "Raqobat strategiyasi", 2)
    doc.add_paragraph(
        "Ko'p profilli markazlar «hamma narsa» taklif qiladi; Radeski esa «teri + soch + tirnoq + "
        "shifokor + zamonaviy texnologiya» tor ixtisoslashuv bilan ajralib turadi. "
        "Bu 5-band tavsiyasiga mos holda saytda qayd etildi."
    )

    # BOB 6
    add_heading(doc, "6-BOB. RADESKI RAQOBAT AFZALLIKLARI", 1)
    doc.add_paragraph("Tahlildagi 8 ta afzallikdan eng muhim 4 tasi saytga kodlandi:")
    add_numbered(doc, [
        "Tor tibbiy ixtisoslashuv — teri, soch va tirnoq",
        "Dermatologiya, trixologiya, podologiya va diagnostika bir joyda",
        "Shifokor qabuli, diagnostika va apparat muolajalari bir rejada",
        "Farg'ona va Qo'qon filiallari — bir xil sifat standarti",
    ])
    add_heading(doc, "Qayerda ko'rinadi", 2)
    add_bullets(doc, [
        "KokandLandingPage.tsx — «Nima uchun Radeski?» bo'limi",
        "FerganaLandingPage.tsx — xuddi shu blok",
        "LocalCommercialLandingPage.tsx — har bir 30 landingda",
    ])

    # BOB 7
    add_heading(doc, "7-BOB. ASOSIY ZAIF JOYLAR VA ULARNI YOPISH", 1)
    doc.add_paragraph("Tahlildagi zaif joylar va amalga oshirilgan yechimlar:")
    weaknesses = [
        ("Har bir commercial so'rov uchun alohida landing yo'q edi", "30 ta alohida URL yaratildi (/uz/qoqon/{slug} va /uz/fargona/{slug})"),
        ("Farg'ona va Qo'qon SEO strategiyasi aralashgan", "Shaharlar ajratildi: alohida hub + alohida 15+15 landing"),
        ("Sitemap to'liq emas edi", "public/sitemap.xml ga 96 ta yangi URL qo'shildi (hreflang bilan)"),
        ("Ichki linking zaif", "Landing → xizmat, condition, promo, maqola, shifokor, natijalar"),
        ("Title/Description shahar aralash", "pageMeta.ts: qoqon va fargona alohida meta; har landing o'z SEO title/desc"),
        ("Canonical va hreflang", "routeSeo.ts: local commercial uchun locale alternates qo'shildi"),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "Muammo (tahlil)"
    table.rows[0].cells[1].text = "Yechim (implementatsiya)"
    for prob, fix in weaknesses:
        row = table.add_row().cells
        row[0].text = prob
        row[1].text = fix

    doc.add_paragraph()
    doc.add_paragraph(
        "Eslatma: shifokorlar shaxsiy SEO (12-band) va 404/redirect monitoring (13-band) "
        "keyingi bosqichda davom ettiriladi — ular 7-bandning qisman yopilgan qismlari."
    )

    # BOB 8
    add_heading(doc, "8-BOB. QO'QON UCHUN SEO YO'NALISHLARI (15 TA SO'ROV)", 1)
    kokand_queries = [
        "Dermatolog Qo'qonda", "Trixolog Qo'qonda", "Podolog Qo'qonda",
        "Akne davolash Qo'qonda", "Soch to'kilishi davolash Qo'qonda",
        "Tirnoq zamburug'i davolash Qo'qonda", "O'sma olib tashlash Qo'qonda",
        "Psoriaz davolash Qo'qonda", "Vitiligo davolash Qo'qonda",
        "Rozasea davolash Qo'qonda", "Pigmentatsiya davolash Qo'qonda",
        "Lazer epilyatsiya Qo'qonda", "Dermatoskopiya Qo'qonda",
        "Qon tomir va qizarish davolash Qo'qonda", "So'gal olib tashlash Qo'qonda",
    ]
    add_table_urls(doc, "qoqon", KOKAND_SLUGS, kokand_queries)
    doc.add_paragraph()
    doc.add_paragraph(
        "Har bir sahifada: H1+shahar, muammo tavsifi, kimlar uchun, davolash usullari, "
        "FAQ (2 ta), qabulga yozilish CTA, bog'liq xizmat/maqola linklari."
    )

    # BOB 9
    add_heading(doc, "9-BOB. FARG'ONA UCHUN PRIORITET SO'ROVLAR (15 TA)", 1)
    fergana_queries = [
        "Dermatolog Farg'onada", "Trixolog Farg'onada", "Podolog Farg'onada",
        "Akne davolash Farg'onada", "Psoriaz davolash Farg'onada",
        "Vitiligo davolash Farg'onada", "Rozasea davolash Farg'onada",
        "Soch to'kilishi Farg'onada", "Dermatoskopiya Farg'onada",
        "O'sma olib tashlash Farg'onada", "Tirnoq zamburug'i Farg'onada",
        "IPL muolajalari Farg'onada", "Pigmentatsiya Farg'onada",
        "Lazer bilan tomir olib tashlash Farg'onada", "Lazer epilyatsiya Farg'onada",
    ]
    add_table_urls(doc, "fargona", FERGANA_SLUGS, fergana_queries)

    # BOB 10
    add_heading(doc, "10-BOB. KONTENT STRATEGIYASI", 1)
    doc.add_paragraph(
        "Asosiy tamoyil: avval commercial landing, keyin maqolalar, shifokor va natijalar bilan qo'llab-quvvatlash."
    )
    add_heading(doc, "Amalga oshirilgan struktura (misol: akne Qo'qon)", 2)
    add_numbered(doc, [
        "Asosiy sahifa: /uz/qoqon/akne-davolash",
        "Xizmat bo'limi: /uz/services/dermatologiya",
        "Kasallik sahifasi: /uz/conditions/acne",
        "Maqolalar: art-akne, art-pryshchi-u-vzroslykh",
        "Natijalar: /uz/results",
        "Shifokorlar: /uz/doctors",
    ])
    add_heading(doc, "Har bir landing talablari (11-band) bajarilishi", 2)
    reqs = [
        "H1 — xizmat + shahar",
        "Muammo tavsifi",
        "Kimlar uchun",
        "Davolash yondashuvi / usullar",
        "Uskunalar (ko'rsatma bo'yicha, masalan Daavlin fototerapiya)",
        "FAQ",
        "Qabulga yozilish CTA",
        "Ichki linklar (xizmat, maqola, shifokor, natija)",
    ]
    add_bullets(doc, reqs)
    doc.add_paragraph(
        "Narx bloki va har bir shifokor profili alohida kuchaytirish — keyingi 30 kunlik "
        "reja (14-band) doirasida davom etadi."
    )

    # Technical summary
    add_heading(doc, "TEXNIK XULOSA", 1)
    add_bullets(doc, [
        "Fayllar: localCommercialSeoCatalog.ts, LocalCommercialLandingPage.tsx, FerganaLandingPage.tsx",
        "Routing: App.tsx, paths.ts (fargona PageId)",
        "SEO: pageMeta.ts, routeSeo.ts (canonical/hreflang)",
        "Sitemap: 300 URL (validate:sitemap OK)",
        "Build: muvaffaqiyatli",
        "Deploy: radeski.uz (git push + VPS rebuild)",
    ])

    add_heading(doc, "KEYINGI QADAMLAR (TAVSIYA)", 1)
    add_numbered(doc, [
        "Google Search Console ga yangi URLlarni yuborish (sitemap.xml)",
        "Har bir landing uchun real shifokor va narx ma'lumotini qo'shish",
        "Maqolalardan commercial sahifalarga anchor-linklar kengaytirish",
        "KPI monitoring: Qo'qon vs Farg'ona alohida (15-band)",
        "Shifokor profillarini kuchaytirish (12-band)",
    ])

    doc.add_paragraph()
    footer = doc.add_paragraph("© Radeski Skin Clinic — SEO implementatsiya hisoboti")
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.save(OUT)
    print(f"Saved: {OUT}")


if __name__ == "__main__":
    main()
