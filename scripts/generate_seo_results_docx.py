#!/usr/bin/env python3
"""Generate Word report: SEO analysis recommendations vs completed work."""
from __future__ import annotations

import re
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
CATALOG = ROOT / "src/data/localCommercialSeoCatalog.ts"
OUT = ROOT / "RADESKI_SEO_TAHLIL_NATIJALARI_07_09_2026.docx"


def slugs_for_city(city: str) -> list[str]:
    text = CATALOG.read_text(encoding="utf-8")
    return re.findall(
        rf"(?:treatmentLanding|specialistLanding)\('{city}', '([^']+)'",
        text,
    )


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_para(doc: Document, text: str, bold: bool = False) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(11)


def add_table_row(table, col0: str, col1: str, col2: str, col3: str) -> None:
    row = table.add_row().cells
    row[0].text = col0
    row[1].text = col1
    row[2].text = col2
    row[3].text = col3


def main() -> None:
    qoqon = slugs_for_city("qoqon")
    fargona = slugs_for_city("fargona")

    doc = Document()
    title = doc.add_heading("ОТЧЁТ О ВЫПОЛНЕНИИ РЕКОМЕНДАЦИЙ", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("Radeski Skin Clinic — radeski.uz")
    r.bold = True
    r.font.size = Pt(12)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(
        f"На основании анализа «анализ с 27.08–03.09.docx»\n"
        f"Дата подготовки отчёта: {date.today().strftime('%d.%m.%Y')}\n"
        f"Статус: работы выполнены и развёрнуты на production (radeski.uz)"
    ).font.size = Pt(10)

    doc.add_paragraph()

    add_heading(doc, "1. Назначение отчёта", 2)
    add_para(
        doc,
        "Данный документ — ответ на SEO-анализ от 27.08–03.09.2026. "
        "По каждому пункту исходного анализа указано: какая была рекомендация / слабое место, "
        "что именно сделано на сайте, статус выполнения и ссылка (если применимо).",
    )

    add_heading(doc, "2. Сводная таблица выполнения", 2)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "№ / Блок анализа"
    hdr[1].text = "Было (рекомендация / слабость)"
    hdr[2].text = "Сделано"
    hdr[3].text = "Статус"

    rows = [
        (
            "4, 8.1",
            "Локальное SEO Коканда слабое (5.5/10). Нужна полноценная страница «Дерматолог в Коканде» / «Dermatolog Qo‘qon» с уникальным контентом.",
            "Hub-страница усилена: SEO-кластеры (заболевания, диагностика, трихология), контакты филиала, блок врачей, 8 экспертных статей, 21 коммерческая landing. UZ/RU/EN.",
            "✅ Выполнено",
        ),
        (
            "5, 9",
            "Фергана и Коканд должны быть разделены: Radeski → город → направление → услуга.",
            f"Отдельные hub: /uz/qoqon и /uz/fargona. Коммерческие страницы: Qo‘qon — {len(qoqon)} шт., Farg‘ona — {len(fargona)} шт. Контент не дублируется.",
            "✅ Выполнено",
        ),
        (
            "8.2 / Задача №2",
            "Страницы заболеваний: акне, псoriasis, витилиgo, розацеа — Фергана и Коканд.",
            "akne-davolash, psoriaz-davolash, vitiligo-davolash, rozasea-davolash — в обоих городах (uz/ru/en).",
            "✅ Выполнено",
        ),
        (
            "6, 8.3 / Задача №3",
            "Диагностика: дерматоскопия, проверка родинок, биопсия, онкодermatология.",
            "dermatoskopiya, xol-tekshiruvi (проверка родинок), biopsiya, onko-dermatolog — в Qo‘qon и Farg‘ona. osma-olib-tashlash — удаление новообразований.",
            "✅ Выполнено",
        ),
        (
            "8.4 / Задача №4",
            "Трихология: трихolog, лечение выпадения волос, трихоскопия.",
            "trixolog, soch-tokilish, trixoskopiya — в обоих городах.",
            "✅ Выполнено",
        ),
        (
            "7.1, 9",
            "Техническая чистота: ошибочные URL, «Процедура не найдена», мусор в индексе.",
            "404 для неверных city-slug; meta robots noindex на ошибках и «процедура не найдена»; sitemap обновлён (132+ URL).",
            "✅ Частично (основное выполнено; полный crawl — следующий этап)",
        ),
        (
            "7.2, 8.7 / Задача №7",
            "Экспертный контент с региональной привязкой, не только SEO-текст.",
            "Мaqola: o‘choqli alopetsiya (Turg‘unov, Qo‘qon). Hub Qo‘qon связан с клиническими статьями. Существующие regional-статьи сохранены.",
            "✅ В процессе (ключевая статья добавлена; серия продолжается)",
        ),
        (
            "3.3, 8.5 / Задача №5",
            "Каждая страница врача — полноценная SEO-посадочная.",
            "Расширены профили: Turg‘unov Shohruz, Usmonova Mohinabonu. На hub Qo‘qon — ссылка на врача филиала.",
            "⚠️ Частично (2 врача; остальные — следующий этап)",
        ),
        (
            "8.6 / Задача №6",
            "5–10 новых подробных отзывов в месяц; Google Maps.",
            "Требует работы клиники / маркетинга, не только разработки сайта.",
            "⏳ Не выполнено (вне scope разработки)",
        ),
        (
            "9",
            "Мониторинг позиций, GSC, страницы которые Google игнорирует.",
            "Sitemap расширен для индексации новых URL; настройка GSC — рекомендуется отдельно.",
            "⏳ Рекомендуется следующий этап",
        ),
    ]
    for row in rows:
        add_table_row(table, *row)

    add_heading(doc, "3. Задача №1 — Коканд (раздел 4 и 8.1 анализа)", 2)
    add_para(doc, "Было:", bold=True)
    add_para(
        doc,
        "«Локальное SEO Коканда — 5.5/10. Нужна полноценная страница под запросы «Дерматолог в Коканде» "
        "и «Dermatolog Qo‘qon» с самостоятельным контентом, врачами филиала, локальными контактами.»",
    )
    add_para(doc, "Сделано:", bold=True)
    bullets_kokand = [
        "Главная hub-страница: https://radeski.uz/uz/qoqon (также /ru/qoqon, /en/qoqon)",
        "H1: «Dermatolog Qo‘qon | Radeski Skin Clinic» / «Дерматолог Коканд»",
        "Уникальный lead-текст с перечислением направлений филиала (не копия Ферганы)",
        "Блок SEO-кластеров: Заболевания | Диагностика | Трихология | Специалисты",
        "Блок врачей филиала (Turg‘unov Shohruz — dermatovenerolog-trixolog)",
        "8 экспертных статей с привязкой к Qo‘qon (акne, psoriaz, vitiligo, alopecia, trixoskopiya и др.)",
        f"21 коммерческая SEO-страница для Qo‘qon (было 15, добавлено 6)",
    ]
    for b in bullets_kokand:
        doc.add_paragraph(b, style="List Bullet")

    add_heading(doc, "4. Задача №2 — Заболевания (раздел 8.2)", 2)
    add_para(doc, "Было: отдельные landing по акne, psoriaz, vitiligo, rozasea для Фergana и Kokand.", bold=True)
    add_para(doc, "Сделано:", bold=True)
    for city, label in [("qoqon", "Qo‘qon"), ("fargona", "Farg‘ona")]:
        doc.add_paragraph(f"{label}:", style="List Bullet")
        for slug in ["akne-davolash", "psoriaz-davolash", "vitiligo-davolash", "rozasea-davolash"]:
            doc.add_paragraph(
                f"https://radeski.uz/uz/{city}/{slug}",
                style="List Bullet 2",
            )

    add_heading(doc, "5. Задача №3 — Диагностика и дерматоонкология (разделы 6 и 8.3)", 2)
    add_para(
        doc,
        "Было: выделить кластер — дерматоскопия, проверка родинок, биопсия, онкодermatolog.",
        bold=True,
    )
    add_para(doc, "Сделано (новые страницы):", bold=True)
    new_diag = [
        ("dermatoskopiya", "Дерматоскопия"),
        ("xol-tekshiruvi", "Проверка родинок (xol/dog tekshiruvi)"),
        ("biopsiya", "Биопсия кожи"),
        ("onko-dermatolog", "Онкодermatolog"),
        ("osma-olib-tashlash", "Удаление новообразований"),
    ]
    for slug, name in new_diag:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(f"{name}: ").bold = True
        p.add_run(f"https://radeski.uz/uz/qoqon/{slug} и https://radeski.uz/uz/fargona/{slug}")

    add_heading(doc, "6. Задача №4 — Трихология (раздел 8.4)", 2)
    add_para(doc, "Было: trixolog, soch to‘kilishi, trixoskopiya — по городам.", bold=True)
    add_para(doc, "Сделано:", bold=True)
    for slug in ["trixolog", "soch-tokilish", "trixoskopiya"]:
        doc.add_paragraph(
            f"https://radeski.uz/uz/qoqon/{slug} | https://radeski.uz/uz/fargona/{slug}",
            style="List Bullet",
        )

    add_heading(doc, "7. Разделение Фergana / Qo‘qon (раздел 5 и 9)", 2)
    add_para(
        doc,
        "Было: «Фергана | Коканд» на общих страницах недостаточно; нужны отдельные центры продвижения.",
        bold=True,
    )
    add_para(doc, "Сделано:", bold=True)
    doc.add_paragraph("Farg‘ona hub: https://radeski.uz/uz/fargona", style="List Bullet")
    doc.add_paragraph("Qo‘qon hub: https://radeski.uz/uz/qoqon", style="List Bullet")
    doc.add_paragraph(
        f"Коммерческие landing: Farg‘ona — {len(fargona)} URL × 3 языка; Qo‘qon — {len(qoqon)} URL × 3 языка",
        style="List Bullet",
    )
    doc.add_paragraph("Перекрёстные ссылки между hub-страницами городов", style="List Bullet")

    add_heading(doc, "8. Техническая чистота (раздел 7.1 и 9)", 2)
    add_para(doc, "Было (проблемы):", bold=True)
    problems = [
        "URL открывает «Процедура не найдена или была удалена»",
        "Ошибочные city-URL падали на hub вместо 404",
        "Sitemap не содержал Farg‘ona hub и коммерческие landing",
        "Ошибочные страницы могли индексироваться Google",
    ]
    for p in problems:
        doc.add_paragraph(p, style="List Bullet")

    add_para(doc, "Сделано:", bold=True)
    fixes = [
        "Неверный slug /uz/qoqon/... или /uz/fargona/... → страница 404 с кнопками «Главная» и «Hub города»",
        "Meta robots noindex, nofollow на: 404 city-slug, «процедура не найдена», «услуга не найдена», неверный promo/condition",
        "public/sitemap.xml: добавлен блок Local commercial SEO — Farg‘ona hub + все landing (132+ URL) с hreflang",
        "Deploy на production VPS (radeski.uz), сборка проверена",
    ]
    for f in fixes:
        doc.add_paragraph(f, style="List Bullet")

    add_heading(doc, "9. Экспертный контент (раздел 7.2, задача №7)", 2)
    add_para(
        doc,
        "Было: переход от SEO-текста к полезным клиническим материалам с региональной привязкой.",
        bold=True,
    )
    add_para(doc, "Сделано:", bold=True)
    articles = [
        "Soch to‘kilishi: diagnostika — klinik holat (Turg‘unov Sh., Qo‘qon) — art-alopecia-areata-klinik-holat",
        "Psoriaz — Daavlin Qo‘qon (существующая regional-статья)",
        "LaseMD Ultra Kokand, Onixokriptoz Kokand, Thulium laser hair Kokand",
        "На hub Qo‘qon добавлены ссылки на 8 экспертных материалов",
    ]
    for a in articles:
        doc.add_paragraph(a, style="List Bullet")

    add_heading(doc, "10. Что НЕ выполнено / следующий этап", 2)
    add_para(
        doc,
        "Следующие пункты исходного анализа требуют действий вне последней разработки или на следующем этапе:",
    )
    pending = [
        "Задача №6: 5–10 новых подробных отзывов в месяц + системная работа с Google Maps (маркетинг клиники)",
        "Задача №5: полные SEO-профили для ВСЕХ врачей (сейчас расширены 2 профиля)",
        "Полный технический crawl всего сайта (Screaming Frog / GSC) — рекомендуется отдельный аудит",
        "Отдельная landing «Mohs-хирургия» — в анализе упомянута, отдельная страница не создавалась",
        "Регулярный мониторинг позиций и Search Console",
    ]
    for p in pending:
        doc.add_paragraph(p, style="List Bullet")

    add_heading(doc, "11. Полный список коммерческих страниц Qo‘qon", 2)
    for slug in qoqon:
        doc.add_paragraph(f"https://radeski.uz/uz/qoqon/{slug}", style="List Bullet")

    add_heading(doc, "12. Полный список коммерческих страниц Farg‘ona", 2)
    for slug in fargona:
        doc.add_paragraph(f"https://radeski.uz/uz/fargona/{slug}", style="List Bullet")

    add_heading(doc, "13. Итог", 2)
    add_para(
        doc,
        "По ключевым рекомендациям SEO-анализа от 27.08–03.09.2026 выполнено: "
        "усиление локального SEO Кokanda, разделение кластеров Farg‘ona/Qo‘qon, "
        "страницы заболеваний, диагностики, трихологии, техническая чистка (404, noindex, sitemap) "
        "и развитие экспертного контента. "
        "Сайт не переделывался заново — усилена существующая архитектура, как рекомендовано в п. 10 исходного анализа. "
        "Отзывы, Google Maps и полные профили всех врачей — следующий управленческий этап.",
    )

    doc.save(OUT)
    print(f"Saved: {OUT}")


if __name__ == "__main__":
    main()
