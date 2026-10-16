#!/usr/bin/env python3
"""Generate Word report on Radeski site development work."""
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Cm

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "RADESKI_ISHLAR_HISOBOTI_2026.docx"


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_bullet(doc: Document, text: str) -> None:
    doc.add_paragraph(text, style="List Bullet")


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        for p in hdr[i].paragraphs:
            for run in p.runs:
                run.bold = True
    for ri, row in enumerate(rows):
        cells = table.rows[ri + 1].cells
        for ci, val in enumerate(row):
            cells[ci].text = val
    doc.add_paragraph()


def main() -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("RADESKI SKIN CLINIC\nSayt rivojlantirish bo‘yicha ish hisoboti")
    run.bold = True
    run.font.size = Pt(16)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run(f"Domen: https://radeski.uz\nSana: {date.today().strftime('%d.%m.%Y')}\n")
    sub.runs[0].font.size = Pt(11)

    doc.add_paragraph()

    add_heading(doc, "1. Qisqacha xulosa", 1)
    doc.add_paragraph(
        "Ushbu hisobot Radeski.uz veb-saytida amalga oshirilgan kontent, funksional, "
        "natijalar (before/after) va SEO-texnik yaxshilashlar bo‘yicha bajarilgan ishlarni "
        "tavsiflaydi. Asosiy maqsad — klinika xizmatlari va maqolalar orqali organik qidiruv "
        "qamrovini kengaytirish, foydalanuvchiga tushunarli tibbiy kontent taqdim etish va "
        "mijozlarni konsultatsiyaga yo‘naltirish."
    )

    add_heading(doc, "2. Yangi maqolalar (kontent)", 1)
    doc.add_paragraph(
        "Barcha yangi maqolalar uch tilda (o‘zbek, rus, ingliz) tayyorlandi: sarlavha, qisqa "
        "tavsif, to‘liq matn, FAQ, teglar, shifokorga murojaat qilish belgilari. "
        "Xizmat sahifalariga ichki havolalar qo‘shildi."
    )

    article_rows = [
        [
            "PhotoFinder diagnostika",
            "photofinder-novoobrazovaniya-diagnostika-radeski",
            "Dermatoskopiya, teri o‘smalari skriningi",
        ],
        [
            "Kunduzgi statsionar",
            "dnevnoj-stacionar-radeski",
            "Kuzatuv, infuziya, ambulator stacionar",
        ],
        [
            "Immunobiologik terapiya",
            "immunobiologik-terapiya-radeski",
            "Surunkali teri kasalliklari",
        ],
        [
            "Lablar kontur plastikasi",
            "konturnaya-plastika-guby-radeski",
            "Inyeksion kosmetologiya",
        ],
        [
            "Soch to‘kilishi profilaktikasi (inyeksiyalar)",
            "soch-tokilishi-profilaktika-inyeksii-radeski",
            "Trixologiya",
        ],
        [
            "Psoriaz — klinik holat (Z. Baxridinov)",
            "psoriaz-klinik-keys-bahridinov-radeski",
            "Psoriaz, UVA 308, fototerapiya",
        ],
        [
            "Tirnoq va oyoq patologiyalari markazi",
            "klinika-patologii-nogtej-stopy-radeski",
            "Podologiya, onixokriptoz",
        ],
        [
            "Klinik tadqiqot (GCP)",
            "klinicheskoe-issledovanie-radeski",
            "Ilmiy faoliyat, klinika",
        ],
        [
            "Kontur plastika (yuz + lablar)",
            "kontur-plastika-gialuron-radeski",
            "Inyeksion kosmetologiya",
        ],
        [
            "Lazer biorevitalizatsiya",
            "lazer-biorevitalizatsiya-radeski",
            "Apparatli kosmetologiya",
        ],
        [
            "Normal va patologik soch to‘kilishi",
            "soch-tokilishi-normal-va-patologiya-radeski",
            "Trixologiya, trixoskopiya (SEO-blok bilan)",
        ],
    ]
    add_table(
        doc,
        ["Mavzu", "URL slug (uz)", "Bog‘liq yo‘nalish"],
        article_rows,
    )
    doc.add_paragraph(
        "Har bir maqola: src/data/articles/*, ro‘yxatdan o‘tkazish (data.ts, articleRichCatalog.ts), "
        "sitemapArticles.ts, serviceRelatedArticles.ts orqali saytga ulandi."
    )

    add_heading(doc, "3. «Natijalar» bo‘limi (davolash natijalari)", 1)
    doc.add_paragraph(
        "Klinik before/after materiallar alohida kartochkalar sifatida qo‘shildi. "
        "Yagona kolaj rasmlar (До|После) comparisonImage rejimida: kartochkada bir xil o‘lcham (2:1), "
        "bosganda lightbox — to‘liq katta ko‘rinish."
    )
    result_rows = [
        ["Vitiligo — yuz repigmentatsiyasi", "1,5 oy", "Faol (kolaj)"],
        ["Psoriaz — plakali o‘choq", "Individual kurs", "Faol"],
        ["CO₂ — operatsiyadan keyingi chandiq (peshona)", "1–3 seans", "Faol"],
        ["CO₂ — postakne shlifovka", "Kurs", "Faol"],
        ["Halqali granulema — 7 kun", "—", "Olib tashlandi (foydalanuvchi so‘rovi)"],
        ["Furunkul — 7 kun", "—", "Olib tashlandi"],
        ["Oyoq tagi so‘g‘alari — CO₂", "—", "Olib tashlandi"],
    ]
    add_table(doc, ["Natija", "Davomiylik", "Holat"], result_rows)
    add_bullet(doc, "Bolalar alopeciya, pustuloz psoriaz va boshqa natijalar avvaldan mavjud; API + statik katalog birlashtiriladi.")
    add_bullet(doc, "Pediatrik natijalarda maxfiylik: ko‘zlar yashirilgan versiyalar.")

    add_heading(doc, "4. Texnik va interfeys o‘zgarishlari", 1)
    for item in [
        "Production build va VPS (161.35.107.0) ga deploy: yangi maqolalar, natijalar, sitemap.",
        "public/sitemap.xml — patchSitemapArticles.ts orqali maqola hreflang bloklari yangilandi.",
        "Natijalar sahifasi: lightbox, filtrlar (xizmat turi bo‘yicha), comparisonImage dizayn sozlamalari.",
        "Navbar: maqolalar va xizmatlar bo‘yicha yaxshilangan dropdown/flyout (mobil moslashuv).",
        "Farg‘ona va Qo‘qon uchun mahalliy SEO landing sahifalari (4–10 bandlar bo‘yicha rejalar).",
        "12 ta dermatologik holat uchun alohida landing (dermatology conditions).",
        "Malaka oshirish / obrazovaniya bo‘limi va alohida dastur sahifalari.",
        "Daavlin uskunalar: model sahifalar, nav integratsiyasi.",
        "Google sharhlari filtrlari va xizmat kategoriyalari bilan bog‘lash.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "5. SEO — nima qilindi va qanday yaxshilanish kutiladi", 1)
    add_heading(doc, "5.1. On-page va texnik SEO (bajarilgan)", 2)
    for item in [
        "Har bir yangi maqola uchun noyob title/description (uz/ru/en); ba’zilarida ARTICLE_SEO_OVERRIDES (maxsus kalit so‘zlar).",
        "Misol: soch to‘kilishi maqolasi — «выпадение волос», «трихолог», «трихоскопия» va hokazo.",
        "articleRichCatalog.ts — qidiruv va rich-snippet uchun teglar, FAQ, keyTakeaways.",
        "Ichki linking: trixologiya, trixoskopiya, PRP, narxlar, tegishli xizmat subpages.",
        "Sitemap: yangi routeKey lar, ustuvorlik (priority) uz/ru/en.",
        "Mahalliy SEO: Farg‘ona/Qo‘qon shahar + xizmat/trixolog/dermatolog landing URL lari.",
        "Strukturali maqola sxemasi (App.tsx — JSON-LD article schema).",
        "Ko‘p tillilik: /uz/, /ru/, /en/ — Google uchun alohida title va description.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "5.2. SEO natijalari (o‘lchov)", 2)
    doc.add_paragraph(
        "Google Search Console va Yandex veb-master panelidan real raqamlar (klik, ko‘rsatish, "
        "o‘rtacha pozitsiya) ushbu hisobotga avtomatik ulanmagan. Quyidagilar mantiqiy kutilma "
        "(2–8 hafta indeksatsiya va kontent yetarli bo‘lgach):"
    )
    for item in [
        "Yangi maqola URL lari indeksatsiyaga kiradi — long-tail so‘rovlar bo‘yicha qamrov o‘sadi.",
        "«Trixolog Farg‘ona/Qo‘qon», «vitiligo davolash», «podologiya», «kontur plastika» kabi so‘rovlar.",
        "Ichki havolalar xizmat sahifalarining muhimligi (internal PageRank) oshadi.",
        "FAQ va rich kontent snippetlar uchun qo‘shimcha imkoniyat.",
        "Natijalar (before/after) — ishonch signali; konversiya va vaqt saytda oshishi mumkin.",
    ]:
        add_bullet(doc, item)
    doc.add_paragraph(
        "Aniq KPI olish uchun tavsiya: GSC da «Performance» — 28 kun solishtirish (ishlar boshlangandan keyin); "
        "maqsadli so‘rovlar bo‘yicha alohida filter; indeksatsiya → «Pages»."
    )

    add_heading(doc, "6. Deploy va fayllar", 1)
    doc.add_paragraph(
        "Server: radeski.uz (VPS), ilova yo‘li /var/www/radeski, build: VITE_API_URL=https://api.radeski.uz. "
        "Deploy skriptlari (scripts/deploy_*_vps.py) orqali maqola va natija yangilanishlari yuklandi."
    )
    add_bullet(doc, "Muhim kataloglar: src/data.ts, sitePagesContent.ts, articleRichCatalog.ts, sitemapArticles.ts, pageMeta.ts")
    add_bullet(doc, "Rasm/media: public/articles/, public/results/, public/services/")

    add_heading(doc, "7. Keyingi qadamlar (tavsiyalar)", 1)
    for item in [
        "GSC / Yandex: yangi URL larni «Inspect URL» → indeksatsiya so‘rovi.",
        "3 oylik SEO monitoring: top 20 so‘rov, CTR, pozitsiya (hisobotni keyingi Word versiyada raqamlar bilan to‘ldirish).",
        "CMS API dagi eski natijalarni admin paneldan tozalash (agar dublikat qolsa).",
        "Maqolalar uchun ijtimoiy tarmoqlarda repost — tezroq trafik.",
        "A/B: natijalar kartochkalarida sarlavha va CTA «Yozilish».",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "8. Xulosa", 1)
    doc.add_paragraph(
        "Davr ichida 11+ yangi to‘liq maqola (3 til), klinik natijalar bo‘limi kengaytirildi, "
        "SEO-texnik infrastruktura (sitemap, meta, ichki havolalar, mahalliy landinglar) mustahkamlangan "
        "va o‘zgarishlar production muhitiga chiqarilgan. SEO natijalari qidiruv tizimlarida "
        "indeksatsiya va reyting o‘sishi shaklida asta-sekin namoyon bo‘ladi; monitoring "
        "Search Console orqali hujjatlashtirish tavsiya etiladi."
    )

    doc.add_paragraph()
    p = doc.add_paragraph("Tayyorladi: loyiha rivojlantirish (Cursor / texnik jamoa)")
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    doc.save(OUT)
    desktop = Path.home() / "Desktop" / OUT.name
    if desktop.parent.exists():
        import shutil

        shutil.copy2(OUT, desktop)
        print(f"Saved: {OUT}")
        print(f"Copy:  {desktop}")
    else:
        print(f"Saved: {OUT}")


if __name__ == "__main__":
    main()
